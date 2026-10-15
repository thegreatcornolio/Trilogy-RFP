#!/usr/bin/env python3
"""Build Trilogy GCC Overview PPTX — Corporate Overview look + Trilogy yellow."""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BLACK = RGBColor(0x0A, 0x0A, 0x0A)
CARD = RGBColor(0x22, 0x22, 0x20)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB8, 0xB8, 0xB4)
YELLOW = RGBColor(0xD5, 0xEC, 0x67)
DARK = RGBColor(0x0A, 0x0A, 0x0A)
LIGHT = RGBColor(0xE6, 0xE6, 0xE6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.55)
FOOTER_H = Inches(0.08)

OUT = "/workspace/GCC/Trilogy-GCC-Overview.pptx"
LOGO = "/workspace/presentations/_shared/logos/trilogy.png"
HERO = "/tmp/gcc-pptx-assets/corp-img-0.png"
CASE_AEGIS = "/workspace/GCC/assets/dbit-case-aegis.png"
CASE_ASC = "/workspace/GCC/assets/dbit-case-ascensos.png"


def set_run(run, size=14, bold=False, color=WHITE):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_text(shape, text, size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return tf


def add_para(tf, text, size=14, bold=False, color=WHITE, space_before=0):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return p


def rect(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def round_rect(slide, left, top, width, height, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def textbox(slide, left, top, width, height, text="", size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(left, top, width, height)
    if text:
        add_text(sh, text, size=size, bold=bold, color=color, align=align)
    return sh


def chrome(slide, page_num, title):
    bg = rect(slide, 0, 0, SLIDE_W, SLIDE_H, BLACK)
    sp_tree = slide.shapes._spTree
    sp = bg._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)

    rect(slide, 0, SLIDE_H - FOOTER_H, SLIDE_W, FOOTER_H, YELLOW)
    textbox(
        slide,
        SLIDE_W - Inches(0.9),
        SLIDE_H - Inches(0.42),
        Inches(0.7),
        Inches(0.3),
        str(page_num),
        size=11,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )
    rect(slide, MARGIN, Inches(0.38), Inches(0.08), Inches(0.38), YELLOW)
    textbox(slide, MARGIN + Inches(0.2), Inches(0.28), Inches(11.5), Inches(0.55), title, size=26, bold=True, color=WHITE)


def footer_note(slide):
    textbox(
        slide,
        MARGIN,
        SLIDE_H - Inches(0.48),
        Inches(7),
        Inches(0.28),
        "Trilogy GCC  ·  Commercial-in-confidence",
        size=9,
        color=MUTED,
    )
    textbox(
        slide,
        SLIDE_W - Inches(4.2),
        SLIDE_H - Inches(0.48),
        Inches(3.2),
        Inches(0.28),
        "proposal.trilogybpo.com/GCC",
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def card(slide, left, top, width, height, fill=CARD):
    return round_rect(slide, left, top, width, height, fill)


def bullets(slide, left, top, width, height, lines, size=13, color=WHITE):
    sh = textbox(slide, left, top, width, height)
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = "•  " + line
        set_run(run, size=size, color=color)
    return sh


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    def new():
        return prs.slides.add_slide(blank)

    n = 0

    # 1 COVER
    n += 1
    s = new()
    rect(s, 0, 0, SLIDE_W, SLIDE_H, BLACK)
    rect(s, 0, 0, Inches(0.12), SLIDE_H, YELLOW)
    rect(s, 0, SLIDE_H - FOOTER_H, SLIDE_W, FOOTER_H, YELLOW)
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, MARGIN, Inches(0.4), height=Inches(0.55))
    textbox(s, MARGIN, Inches(1.35), Inches(7.2), Inches(0.35), "GCC OVERVIEW", size=12, bold=True, color=YELLOW)
    title = textbox(s, MARGIN, Inches(1.75), Inches(7.5), Inches(1.5))
    tf = title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Global Capability Centre"
    set_run(r, size=36, bold=True, color=WHITE)
    add_para(tf, "Design, Build, Operate, Transfer", size=28, bold=True, color=YELLOW, space_before=8)
    body = textbox(s, MARGIN, Inches(3.55), Inches(7.2), Inches(1.5))
    tf = body.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = (
        "Mitigating the risks of building a captive centre offshore — how Trilogy delivers DBOT, "
        "with a transparent costing model and a fixed management fee."
    )
    set_run(r, size=15, color=MUTED)
    textbox(s, MARGIN, Inches(5.35), Inches(7), Inches(0.3), "Human Led  ·  Ai-enabled", size=13, bold=True, color=WHITE)
    textbox(s, MARGIN, Inches(5.75), Inches(7), Inches(0.3), "Trilogy GCC  ·  A Trilogy Group Company", size=12, color=MUTED)
    textbox(s, MARGIN, Inches(6.15), Inches(7), Inches(0.3), "sales@trilogybpo.com  ·  trilogybpo.com", size=12, color=YELLOW)
    textbox(
        s,
        SLIDE_W - Inches(0.9),
        SLIDE_H - Inches(0.42),
        Inches(0.7),
        Inches(0.3),
        str(n),
        size=11,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )
    if os.path.exists(HERO):
        s.shapes.add_picture(HERO, Inches(8.55), Inches(0.55), height=Inches(6.35))

    # 2 AGENDA
    n += 1
    s = new()
    chrome(s, n, "Agenda")
    items = [
        ("01", "Company overview — Trilogy GCC"),
        ("02", "Why South Africa"),
        ("03", "Three routes offshore"),
        ("04", "Business risks of a captive"),
        ("05", "The solution — DBOT"),
        ("06", "Implementation phases"),
        ("07", "Benefits & commercial model"),
        ("08", "Case studies & next steps"),
    ]
    for i, (num, label) in enumerate(items):
        col, row = divmod(i, 4)
        x = MARGIN + col * Inches(6.2)
        y = Inches(1.15) + row * Inches(1.2)
        card(s, x, y, Inches(5.85), Inches(1.05), CARD)
        textbox(s, x + Inches(0.25), y + Inches(0.28), Inches(0.7), Inches(0.45), num, size=22, bold=True, color=YELLOW)
        textbox(s, x + Inches(1.1), y + Inches(0.32), Inches(4.4), Inches(0.45), label, size=15, bold=True, color=WHITE)
    footer_note(s)

    # 3 COMPANY OVERVIEW
    n += 1
    s = new()
    chrome(s, n, "01  —  Company overview")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.55),
        "Move beyond outsourcing. Establish a Global Capability Centre in South Africa that functions as a seamless extension of your headquarters.",
        size=14,
        color=MUTED,
    )
    pills = [
        ("Build your own centre", "Design, Build, Operate, Transfer — a risk-mitigated path to a Center of Excellence in South Africa."),
        ("20–30% vs BPO", "Global capability centres / captives deliver 20–30% cost savings versus offshore outsourcing."),
        ("Digital & autonomous", "Unique digital and autonomous platforms to reduce headcount while improving CX."),
        ("1,000+ FTEs", "Owned and operated footprint across Cape Town and Johannesburg — never franchise."),
    ]
    for i, (h, body) in enumerate(pills):
        x = MARGIN + (i % 2) * Inches(6.2)
        y = Inches(1.7) + (i // 2) * Inches(2.25)
        card(s, x, y, Inches(5.95), Inches(2.05), CARD)
        rect(s, x, y, Inches(0.1), Inches(2.05), YELLOW)
        textbox(s, x + Inches(0.35), y + Inches(0.35), Inches(5.3), Inches(0.4), h, size=18, bold=True, color=YELLOW)
        textbox(s, x + Inches(0.35), y + Inches(0.9), Inches(5.3), Inches(0.9), body, size=13, color=WHITE)
    footer_note(s)

    # 4 ACHIEVEMENTS
    n += 1
    s = new()
    chrome(s, n, "Our achievements")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(8),
        Inches(0.7),
        "Trilogy's 20-year track record speaks to our ability to deliver results at scale — creating jobs, building infrastructure, and pioneering technology across the BPO industry.",
        size=13,
        color=MUTED,
    )
    stats = [
        ("45K", "Jobs created", "Significant positive impact on South Africa's economy and communities over 20 years."),
        ("20+", "Years of experience", "Two decades of proven value delivery across multiple industries and geographies."),
        ("8", "Years in AI / ML", "Developing AI, ML platforms, analytics solutions, and digital automation tools."),
        ("36", "Blue-chip centres built", "Contact centres successfully built and launched for blue-chip companies across SA."),
    ]
    for i, (val, label, desc) in enumerate(stats):
        x = MARGIN + (i % 2) * Inches(4.0)
        y = Inches(1.85) + (i // 2) * Inches(2.2)
        textbox(s, x, y, Inches(3.7), Inches(0.55), val, size=36, bold=True, color=YELLOW)
        textbox(s, x, y + Inches(0.55), Inches(3.7), Inches(0.35), label, size=16, bold=True, color=WHITE)
        textbox(s, x, y + Inches(1.0), Inches(3.7), Inches(0.9), desc, size=12, color=MUTED)
    if os.path.exists(HERO):
        s.shapes.add_picture(HERO, Inches(8.7), Inches(1.1), height=Inches(5.5))
    footer_note(s)

    # 5 WHY SA
    n += 1
    s = new()
    chrome(s, n, "Why South Africa?")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.55),
        "South Africa has consistently been ranked as the second most attractive BPO location globally. Source: McKinsey",
        size=13,
        color=MUTED,
    )
    card(s, MARGIN, Inches(1.65), Inches(6.0), Inches(1.55), YELLOW)
    textbox(s, MARGIN + Inches(0.3), Inches(1.8), Inches(5.4), Inches(0.35), "Cost efficiency meets quality", size=15, bold=True, color=DARK)
    textbox(
        s,
        MARGIN + Inches(0.3),
        Inches(2.2),
        Inches(5.4),
        Inches(0.8),
        "Up to 60% lower operational costs than UK/USA markets — wage arbitrage without compromising quality.",
        size=12,
        color=DARK,
    )
    card(s, MARGIN + Inches(6.25), Inches(1.65), Inches(6.0), Inches(1.55), LIGHT)
    textbox(s, MARGIN + Inches(6.55), Inches(1.8), Inches(5.4), Inches(0.35), "Cape Town: the premium hub", size=15, bold=True, color=DARK)
    textbox(
        s,
        MARGIN + Inches(6.55),
        Inches(2.2),
        Inches(5.4),
        Inches(0.8),
        "Leading BPO destination with rapid growth, exceptional infrastructure, and the largest skilled talent pool in the region.",
        size=12,
        color=DARK,
    )
    mid = [
        ("#2 globally", "Most attractive BPO location (McKinsey)"),
        ("GMT+2", "1–2 hrs from UK · extended US coverage"),
        ("95%+", "English proficiency with neutral accents"),
        ("Top-3", "Global BPO destination for three consecutive years"),
        ("Young talent", "Educated, professional workforce ready to deliver"),
        ("Full CX stack", "Voice, tech support, digital and advanced CX"),
    ]
    for i, (h, b) in enumerate(mid):
        x = MARGIN + (i % 3) * Inches(4.15)
        y = Inches(3.5) + (i // 3) * Inches(1.55)
        card(s, x, y, Inches(3.95), Inches(1.4), CARD)
        textbox(s, x + Inches(0.25), y + Inches(0.25), Inches(3.45), Inches(0.35), h, size=16, bold=True, color=YELLOW)
        textbox(s, x + Inches(0.25), y + Inches(0.7), Inches(3.45), Inches(0.5), b, size=12, color=WHITE)
    footer_note(s)

    # 6 THREE ROUTES
    n += 1
    s = new()
    chrome(s, n, "Three routes offshore")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.55),
        "Companies can establish an offshore capability centre in South Africa — for any type of operation — in one of three ways.",
        size=13,
        color=MUTED,
    )
    routes = [
        ("1", "Outsourcing", "Outsource to a competent and credible partner in South Africa under an outsourcing contract.", False),
        ("2", "Captive, own resources", "Establish a captive centre in South Africa with your own resources — full ownership, full localisation risk.", False),
        ("3", "Captive with local partner — DBOT", "Establish a captive with a local partner as a phased Design, Build, Operate and Transfer approach.", True),
    ]
    for i, (num, title, body, focus) in enumerate(routes):
        x = MARGIN + i * Inches(4.15)
        fill = YELLOW if focus else CARD
        tcol = DARK if focus else WHITE
        bcol = DARK if focus else MUTED
        card(s, x, Inches(1.75), Inches(3.95), Inches(4.4), fill)
        textbox(
            s,
            x + Inches(0.3),
            Inches(2.05),
            Inches(3.3),
            Inches(0.45),
            num,
            size=28,
            bold=True,
            color=(DARK if focus else YELLOW),
        )
        textbox(s, x + Inches(0.3), Inches(2.65), Inches(3.3), Inches(0.9), title, size=18, bold=True, color=tcol)
        textbox(s, x + Inches(0.3), Inches(3.7), Inches(3.3), Inches(1.8), body, size=13, color=bcol)
        if focus:
            textbox(s, x + Inches(0.3), Inches(5.55), Inches(3.3), Inches(0.35), "FOCUS OF THIS BRIEF", size=11, bold=True, color=DARK)
    footer_note(s)

    # 7 FIVE RISKS
    n += 1
    s = new()
    chrome(s, n, "The business case — five risk areas")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.55),
        "Establishing a captive offshore has clear financial benefits — and inherent risks that can destroy the value proposition. DBOT is designed to mitigate all five.",
        size=13,
        color=MUTED,
    )
    risks = [
        ("01", "Localisation", "Labour law, equity, location, culture, protocol and contractual pitfalls."),
        ("02", "Financial", "CAPEX for infrastructure, IT, telephony, connectivity and facilities — and shortened ROI pressure."),
        ("03", "Business", "Defocusing management; standing up HR, finance, facilities, legal and ops IT offshore."),
        ("04", "Implementation", "Speed to go-live, scope creep, travel, recruitment and training costs."),
        ("05", "Ongoing ops", "Performance management, fix-it travel, transferring onshore practices to local management."),
    ]
    for i, (num, title, body) in enumerate(risks):
        x = MARGIN + i * Inches(2.48)
        card(s, x, Inches(1.75), Inches(2.35), Inches(4.4), CARD)
        rect(s, x, Inches(1.75), Inches(2.35), Inches(0.12), YELLOW)
        textbox(s, x + Inches(0.18), Inches(2.1), Inches(2.0), Inches(0.4), num, size=20, bold=True, color=YELLOW)
        textbox(s, x + Inches(0.18), Inches(2.6), Inches(2.0), Inches(0.7), title, size=15, bold=True, color=WHITE)
        textbox(s, x + Inches(0.18), Inches(3.4), Inches(2.0), Inches(2.4), body, size=12, color=MUTED)
    footer_note(s)

    # 8 SOLUTION DBOT
    n += 1
    s = new()
    chrome(s, n, "The solution — DBOT hybrid")
    # equation row
    card(s, MARGIN, Inches(1.1), Inches(2.9), Inches(0.65), CARD)
    textbox(s, MARGIN + Inches(0.1), Inches(1.22), Inches(2.7), Inches(0.4), "Outsourcing", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    textbox(s, MARGIN + Inches(2.95), Inches(1.2), Inches(0.4), Inches(0.45), "+", size=22, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
    card(s, MARGIN + Inches(3.4), Inches(1.1), Inches(4.3), Inches(0.65), CARD)
    textbox(
        s,
        MARGIN + Inches(3.5),
        Inches(1.22),
        Inches(4.1),
        Inches(0.4),
        "Captive on internal resources",
        size=14,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    textbox(s, MARGIN + Inches(7.75), Inches(1.2), Inches(0.4), Inches(0.45), "=", size=22, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)
    card(s, MARGIN + Inches(8.2), Inches(1.1), Inches(4.0), Inches(0.65), YELLOW)
    textbox(
        s,
        MARGIN + Inches(8.3),
        Inches(1.22),
        Inches(3.8),
        Inches(0.4),
        "DBOT hybrid",
        size=14,
        bold=True,
        color=DARK,
        align=PP_ALIGN.CENTER,
    )

    textbox(
        s,
        MARGIN,
        Inches(2.0),
        Inches(12.2),
        Inches(0.65),
        "DBOT — Design, Build, Operate and Transfer — is a phased turnkey approach with a highly qualified team of internal and local experts, mitigating captive risk while keeping a path to ownership.",
        size=13,
        color=MUTED,
    )
    phases = [
        ("D", "Design", "Due diligence, gap analysis, solution design, implementation, change and project plans."),
        ("B", "Build", "Total solution implemented at the location of choice, governed by a project team."),
        ("O", "Operate", "Full operational strategy run for a defined period to stabilise operations."),
        ("T", "Transfer", "Operations, technology, facilities, people, knowledge and IP transferred to you."),
    ]
    for i, (letter, name, desc) in enumerate(phases):
        x = MARGIN + i * Inches(3.15)
        card(s, x, Inches(2.85), Inches(3.0), Inches(3.5), CARD)
        rect(s, x, Inches(2.85), Inches(3.0), Inches(0.85), YELLOW if i < 3 else WHITE)
        textbox(s, x + Inches(0.2), Inches(2.95), Inches(0.5), Inches(0.6), letter, size=26, bold=True, color=DARK)
        textbox(s, x + Inches(0.7), Inches(3.1), Inches(2.1), Inches(0.4), name.upper(), size=14, bold=True, color=DARK)
        textbox(s, x + Inches(0.2), Inches(3.95), Inches(2.6), Inches(2.1), desc, size=13, color=WHITE)
    footer_note(s)

    # 9 PHASE DETAIL
    n += 1
    s = new()
    chrome(s, n, "Implementation — phase by phase")
    details = [
        ("Design", ["Comprehensive due diligence", "Local practices gap analysis", "Optimal solution design", "Implementation & change plans", "Project management plan"]),
        ("Build", ["Site & IT build-out (turnkey)", "Agent lifecycle framework", "Recruitment, training, incubation", "Workforce optimisation & QA", "Go-live readiness gates"]),
        ("Operate", ["Full operational strategy", "Integrated management system", "KPI / SLA measurement", "Governed through an MSA", "Stabilise quality & service levels"]),
        ("Transfer", ["People & knowledge transfer", "Technology, facilities & IP", "Flexible timing (e.g. 24/36 mo)", "Optional ongoing support", "Overflow · BC/DR · benchmarking"]),
    ]
    for i, (title, lines) in enumerate(details):
        x = MARGIN + i * Inches(3.15)
        card(s, x, Inches(1.1), Inches(3.0), Inches(5.4), CARD)
        rect(s, x, Inches(1.1), Inches(3.0), Inches(0.1), YELLOW)
        textbox(s, x + Inches(0.2), Inches(1.35), Inches(2.6), Inches(0.4), title.upper(), size=16, bold=True, color=YELLOW)
        bullets(s, x + Inches(0.2), Inches(1.9), Inches(2.6), Inches(4.3), lines, size=12, color=WHITE)
    footer_note(s)

    # 10 BENEFITS
    n += 1
    s = new()
    chrome(s, n, "Business benefits — how DBOT mitigates risk")
    benefits = [
        "Risk-averse localisation model",
        "Limited CAPEX — or an option for no CAPEX",
        "Speed to offshore markets and implementation",
        "Speed to the cost savings anticipated",
        "Instant access to local admin, legal and financial resources",
        "Instant access to local business practices and infrastructure",
        "Mitigates labour legislation, unions and performance-management risk",
        "Aligned incentives — fixed management fee, not headcount tax",
    ]
    for i, b in enumerate(benefits):
        x = MARGIN + (i % 2) * Inches(6.2)
        y = Inches(1.1) + (i // 2) * Inches(1.2)
        card(s, x, y, Inches(5.95), Inches(1.05), CARD)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.25), y + Inches(0.3), Inches(0.4), Inches(0.4))
        circ.fill.solid()
        circ.fill.fore_color.rgb = YELLOW
        circ.line.fill.background()
        textbox(
            s,
            x + Inches(0.3),
            y + Inches(0.35),
            Inches(0.35),
            Inches(0.3),
            str(i + 1),
            size=12,
            bold=True,
            color=DARK,
            align=PP_ALIGN.CENTER,
        )
        textbox(s, x + Inches(0.85), y + Inches(0.32), Inches(4.85), Inches(0.5), b, size=14, bold=True, color=WHITE)
    footer_note(s)

    # 11 COMMERCIAL
    n += 1
    s = new()
    chrome(s, n, "Commercial model")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.45),
        "Transparent costing by DBOT phase, plus a fixed management fee — always with a path to ownership.",
        size=13,
        color=MUTED,
    )
    card(s, MARGIN, Inches(1.55), Inches(6.0), Inches(4.7), CARD)
    textbox(s, MARGIN + Inches(0.35), Inches(1.75), Inches(5.3), Inches(0.4), "Transparent costing + fixed fee", size=18, bold=True, color=YELLOW)
    bullets(
        s,
        MARGIN + Inches(0.35),
        Inches(2.35),
        Inches(5.3),
        Inches(3.6),
        [
            "Costed openly by Design / Build / Operate / Transfer",
            "Transfer typically at 24 or 36 months",
            "Trilogy management fee is fixed — decoupled from headcount",
            "You keep the savings from efficiency and automation",
            "Assumptive model sized on seats, bodies, ratios and location",
        ],
        size=13,
    )
    card(s, MARGIN + Inches(6.25), Inches(1.55), Inches(6.0), Inches(2.2), CARD)
    textbox(s, MARGIN + Inches(6.6), Inches(1.75), Inches(5.4), Inches(0.35), "CAPEX — green-field build", size=16, bold=True, color=WHITE)
    textbox(
        s,
        MARGIN + Inches(6.6),
        Inches(2.25),
        Inches(5.4),
        Inches(1.2),
        "Customer meets the CAPEX for site, equipment and facilities. Full ownership of assets from day one.",
        size=13,
        color=MUTED,
    )
    card(s, MARGIN + Inches(6.25), Inches(3.95), Inches(6.0), Inches(2.3), YELLOW)
    textbox(s, MARGIN + Inches(6.6), Inches(4.15), Inches(5.4), Inches(0.35), "OPEX — Trilogy rent-to-own", size=16, bold=True, color=DARK)
    textbox(
        s,
        MARGIN + Inches(6.6),
        Inches(4.65),
        Inches(5.4),
        Inches(1.3),
        "Full bill of materials on a 36-month payment plan. Site, equipment and facilities are owned by the client at month 36. Chosen by a number of our GCC customers.",
        size=13,
        color=DARK,
    )
    footer_note(s)

    # 12 THREE OFFERINGS
    n += 1
    s = new()
    chrome(s, n, "Where GCC sits in Trilogy's offerings")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.55),
        "Three pathways to elevate CX — outsourcing, owning a GCC, or AI/digital platforms. GCC delivers the additional 20–30% beyond BPO labour arbitrage.",
        size=13,
        color=MUTED,
    )
    offers = [
        ("Offshoring BPO", "Voice to Digital to Autonomous CX", "Outsource to Trilogy in South Africa with modern platforms ready — just add humans.", "40–60% saving vs captive UK / outsourced UK"),
        ("Global Capability Centres", "Design, Build, Operate & Transfer", "Build your own capability centre with Trilogy's DBOT methodology and local expertise.", "Additional 20–30% vs outsourced SA BPO"),
        ("AI & Digital Platforms", "Human in the loop", "Award-winning AI-enabled digital and autonomous platforms for service, sales and collections.", "Further savings through digital & AI deflection"),
    ]
    for i, (title, sub, body, foot) in enumerate(offers):
        x = MARGIN + i * Inches(4.15)
        card(s, x, Inches(1.7), Inches(3.95), Inches(4.55), CARD)
        textbox(s, x + Inches(0.25), Inches(1.95), Inches(3.45), Inches(0.7), title, size=16, bold=True, color=WHITE)
        textbox(s, x + Inches(0.25), Inches(2.7), Inches(3.45), Inches(0.4), sub, size=12, bold=True, color=YELLOW)
        textbox(s, x + Inches(0.25), Inches(3.25), Inches(3.45), Inches(1.6), body, size=12, color=MUTED)
        rect(s, x, Inches(5.35), Inches(3.95), Inches(0.9), YELLOW)
        textbox(s, x + Inches(0.2), Inches(5.5), Inches(3.55), Inches(0.65), foot, size=11, bold=True, color=DARK)
    footer_note(s)

    # 13 CASE AEGIS
    n += 1
    s = new()
    chrome(s, n, "Case study — Aegis South Africa (now Startek)")
    if os.path.exists(CASE_AEGIS):
        s.shapes.add_picture(CASE_AEGIS, MARGIN, Inches(1.15), width=Inches(6.4))
    bullets(
        s,
        Inches(7.2),
        Inches(1.3),
        Inches(5.5),
        Inches(5.0),
        [
            "Designed, built and operated in two SA locations",
            "Grew by 2,500 seats within 5 years",
            "EBITDA of over 15%",
            "Brand still operating successfully as Startek",
            "Proof that DBOT can scale a lasting in-country operation",
        ],
        size=15,
    )
    footer_note(s)

    # 14 CASE ASCENSOS
    n += 1
    s = new()
    chrome(s, n, "Case study — Ascensos")
    if os.path.exists(CASE_ASC):
        s.shapes.add_picture(CASE_ASC, MARGIN, Inches(1.15), width=Inches(6.4))
    bullets(
        s,
        Inches(7.2),
        Inches(1.3),
        Inches(5.5),
        Inches(5.0),
        [
            "First 100 FTEs live in Cape Town within six weeks",
            "1,500 seats over 3 years across 4 locations",
            "$6M+ in government incentives secured",
            "Designed and built a centre of excellence",
            "Operated 2–4 years, then handed over a world-class operation",
        ],
        size=15,
    )
    footer_note(s)

    # 15 LOCATIONS
    n += 1
    s = new()
    chrome(s, n, "South African footprint")
    textbox(
        s,
        MARGIN,
        Inches(0.95),
        Inches(12.2),
        Inches(0.55),
        "Every site is Trilogy owned or leased and operated directly — never via franchise or affiliate. Your operation fits seamlessly into our current Cape Town base.",
        size=13,
        color=MUTED,
    )
    locs = [
        ("Century City", "Cape Town", "300 seats"),
        ("Head Office — Pinelands", "Cape Town", "1,500 seats"),
        ("Newtown", "Johannesburg", "400 seats"),
        ("Rosebank", "Johannesburg", "250 seats"),
    ]
    for i, (name, city, seats) in enumerate(locs):
        x = MARGIN + i * Inches(3.15)
        card(s, x, Inches(1.8), Inches(3.0), Inches(3.6), CARD)
        rect(s, x, Inches(1.8), Inches(3.0), Inches(0.12), YELLOW)
        textbox(s, x + Inches(0.25), Inches(2.3), Inches(2.5), Inches(1.0), name, size=18, bold=True, color=WHITE)
        textbox(s, x + Inches(0.25), Inches(3.4), Inches(2.5), Inches(0.4), city, size=13, color=MUTED)
        textbox(s, x + Inches(0.25), Inches(4.2), Inches(2.5), Inches(0.5), seats, size=22, bold=True, color=YELLOW)
    footer_note(s)

    # 16 CLOSE
    n += 1
    s = new()
    rect(s, 0, 0, SLIDE_W, SLIDE_H, BLACK)
    rect(s, 0, 0, Inches(0.12), SLIDE_H, YELLOW)
    rect(s, 0, SLIDE_H - FOOTER_H, SLIDE_W, FOOTER_H, YELLOW)
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, MARGIN, Inches(0.45), height=Inches(0.55))
    textbox(s, MARGIN, Inches(2.0), Inches(12), Inches(0.8), "Let's build your GCC", size=36, bold=True, color=WHITE)
    textbox(
        s,
        MARGIN,
        Inches(2.9),
        Inches(11),
        Inches(0.8),
        "Design, Build, Operate, Transfer — from South Africa for UK, US and global brands.",
        size=18,
        color=YELLOW,
    )
    textbox(s, MARGIN, Inches(4.1), Inches(11), Inches(0.4), "sales@trilogybpo.com", size=20, bold=True, color=WHITE)
    textbox(s, MARGIN, Inches(4.65), Inches(11), Inches(0.35), "trilogybpo.com   ·   proposal.trilogybpo.com/GCC", size=14, color=MUTED)
    textbox(s, MARGIN, Inches(5.4), Inches(11), Inches(0.35), "Trilogy CX (Pty) Ltd  ·  Commercial-in-confidence", size=12, color=MUTED)
    textbox(
        s,
        SLIDE_W - Inches(0.9),
        SLIDE_H - Inches(0.42),
        Inches(0.7),
        Inches(0.3),
        str(n),
        size=11,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )

    prs.save(OUT)
    print(f"Saved {OUT} with {n} slides ({os.path.getsize(OUT)} bytes)")


if __name__ == "__main__":
    build()
